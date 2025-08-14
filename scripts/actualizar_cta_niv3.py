#!/usr/bin/env python3
"""
Actualizador de cta_niv3 - Plan de Cuentas Contables SISK13 (KHIPUS)
=====================================================================

Este script actualiza la columna cta_niv3 del plan de cuentas según las reglas:
- Nivel 1 y 2: cta_niv3 debe estar vacío
- Nivel 3+: cta_niv3 = primeros 3 dígitos de cuenta + 7 ceros

Las celdas modificadas se resaltan con fondo amarillo.

Uso:
    python actualizar_cta_niv3.py [ruta_archivo_excel]

Autor: Claude Code
Fecha: 2025-01-12
"""

import pandas as pd
import sys
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import shutil
import argparse

class ActualizadorCtaNiv3:
    def __init__(self, archivo_excel=None):
        self.archivo_excel = archivo_excel or 'plan_ctas/plan_cuentas.xlsx'
        self.df = None
        self.cambios = []
        self.workbook = None
        self.worksheet = None
        
    def cargar_archivo(self):
        """Carga el archivo Excel del plan de cuentas"""
        try:
            if not os.path.exists(self.archivo_excel):
                raise FileNotFoundError(f"No se encontro el archivo: {self.archivo_excel}")
                
            print(f"Cargando archivo: {self.archivo_excel}")
            self.df = pd.read_excel(self.archivo_excel)
            
            # Filtrar solo filas con datos validos
            self.df = self.df[self.df['cuenta'].notna()].copy()
            
            print(f"Archivo cargado exitosamente - {len(self.df)} cuentas encontradas")
            
            # Verificar columnas requeridas
            columnas_requeridas = ['cuenta', 'cn_nivel', 'cta_niv3']
            for col in columnas_requeridas:
                if col not in self.df.columns:
                    raise ValueError(f"El archivo debe contener la columna '{col}'")
                    
            return True
            
        except Exception as e:
            print(f"ERROR al cargar el archivo: {e}")
            return False
    
    def calcular_cta_niv3_correcto(self, cuenta_str, nivel):
        """Calcula el valor correcto para cta_niv3"""
        if nivel <= 2:
            return None  # Debe estar vacio
        else:
            cuenta = str(int(float(cuenta_str))).zfill(10)
            return int(cuenta[:3] + '0000000')
    
    def analizar_cambios_necesarios(self):
        """Analiza qué cambios son necesarios"""
        print("\nAnalizando cambios necesarios...")
        
        self.cambios = []
        
        for i, row in self.df.iterrows():
            cuenta_num = row['cuenta']
            nivel = int(row['cn_nivel']) if pd.notna(row['cn_nivel']) else 0
            cta_niv3_actual = row['cta_niv3'] if pd.notna(row['cta_niv3']) else None
            
            # Calcular valor correcto
            cta_niv3_correcto = self.calcular_cta_niv3_correcto(cuenta_num, nivel)
            
            # Comparar actual vs correcto
            necesita_cambio = False
            
            if cta_niv3_correcto is None:  # Debe estar vacio
                if pd.notna(cta_niv3_actual):
                    necesita_cambio = True
            else:  # Debe tener valor especifico
                if pd.isna(cta_niv3_actual) or int(cta_niv3_actual) != cta_niv3_correcto:
                    necesita_cambio = True
            
            if necesita_cambio:
                cuenta_str = str(int(cuenta_num)).zfill(10)
                actual_str = str(int(cta_niv3_actual)).zfill(10) if pd.notna(cta_niv3_actual) else 'VACIO'
                correcto_str = str(cta_niv3_correcto).zfill(10) if cta_niv3_correcto is not None else 'VACIO'
                
                self.cambios.append({
                    'fila_excel': i + 2,  # +2 porque Excel cuenta desde 1 y hay header
                    'fila_pandas': i,
                    'cuenta': cuenta_str,
                    'nivel': nivel,
                    'descripcion': row['descri'] if pd.notna(row['descri']) else 'Sin descripcion',
                    'actual': cta_niv3_actual,
                    'correcto': cta_niv3_correcto,
                    'actual_str': actual_str,
                    'correcto_str': correcto_str
                })
        
        print(f"Cambios necesarios identificados: {len(self.cambios)}")
        return len(self.cambios) > 0
    
    def mostrar_cambios_pendientes(self):
        """Muestra los cambios que se van a realizar"""
        if not self.cambios:
            print("\nNo hay cambios necesarios. Todas las cuentas tienen cta_niv3 correcto.")
            return
        
        print(f"\n{'='*70}")
        print(f"CAMBIOS A REALIZAR EN CTA_NIV3")
        print(f"{'='*70}")
        
        print(f"{'Fila':<6} {'Cuenta':<12} {'Nivel':<6} {'Actual':<12} {'Correcto':<12} {'Descripcion'}")
        print("-" * 70)
        
        for cambio in self.cambios[:20]:  # Mostrar maximo 20
            print(f"{cambio['fila_excel']:<6} {cambio['cuenta']:<12} {cambio['nivel']:<6} "
                  f"{cambio['actual_str']:<12} {cambio['correcto_str']:<12} "
                  f"{cambio['descripcion'][:20]}...")
        
        if len(self.cambios) > 20:
            print(f"\n... y {len(self.cambios) - 20} cambios mas.")
        
        print(f"\nTOTAL DE CAMBIOS: {len(self.cambios)}")
    
    def crear_backup(self):
        """Crea una copia de seguridad del archivo original"""
        try:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            backup_filename = f"{self.archivo_excel.rsplit('.', 1)[0]}_backup_{timestamp}.xlsx"
            
            shutil.copy2(self.archivo_excel, backup_filename)
            print(f"Backup creado: {backup_filename}")
            return backup_filename
            
        except Exception as e:
            print(f"ERROR al crear backup: {e}")
            return None
    
    def aplicar_cambios(self):
        """Aplica los cambios al archivo Excel con formato amarillo"""
        if not self.cambios:
            print("No hay cambios que aplicar.")
            return True
            
        try:
            print("\nAplicando cambios al archivo Excel...")
            
            # Cargar el workbook para manipular formatos
            self.workbook = load_workbook(self.archivo_excel)
            self.worksheet = self.workbook.active
            
            # Buscar la columna cta_niv3
            cta_niv3_col = None
            for col in range(1, self.worksheet.max_column + 1):
                if self.worksheet.cell(1, col).value == 'cta_niv3':
                    cta_niv3_col = col
                    break
            
            if cta_niv3_col is None:
                raise ValueError("No se encontro la columna 'cta_niv3' en el Excel")
            
            # Definir formato amarillo
            fill_amarillo = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
            
            # Aplicar cambios
            cambios_aplicados = 0
            for cambio in self.cambios:
                fila_excel = cambio['fila_excel']
                valor_nuevo = cambio['correcto']
                
                # Actualizar el DataFrame
                self.df.loc[cambio['fila_pandas'], 'cta_niv3'] = valor_nuevo
                
                # Actualizar la celda en Excel y aplicar formato
                celda = self.worksheet.cell(fila_excel, cta_niv3_col)
                celda.value = valor_nuevo
                celda.fill = fill_amarillo
                
                cambios_aplicados += 1
            
            print(f"Cambios aplicados: {cambios_aplicados}")
            return True
            
        except Exception as e:
            print(f"ERROR al aplicar cambios: {e}")
            return False
    
    def guardar_archivo(self):
        """Guarda el archivo con los cambios"""
        try:
            print("Guardando archivo...")
            self.workbook.save(self.archivo_excel)
            print("Archivo guardado exitosamente.")
            return True
            
        except Exception as e:
            print(f"ERROR al guardar archivo: {e}")
            return False
    
    def generar_sql_actualizacion(self, archivo_sql=None):
        """Genera archivo SQL con las actualizaciones necesarias para cta_niv3"""
        if archivo_sql is None:
            archivo_sql = 'plan_ctas/update_cta_niv3.sql'
        
        try:
            with open(archivo_sql, 'w', encoding='utf-8') as f:
                f.write(f"-- ACTUALIZACION DE CTA_NIV3 - PLAN DE CUENTAS CONTABLES\n")
                f.write(f"-- ======================================================\n")
                f.write(f"-- Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"-- Archivo origen: {self.archivo_excel}\n")
                f.write(f"-- Tabla destino: arcgms\n")
                f.write(f"-- Total de cuentas: {len(self.df)}\n")
                f.write(f"-- Cambios necesarios: {len(self.cambios)}\n\n")
                
                if self.cambios:
                    f.write(f"-- ACTUALIZACIONES NECESARIAS:\n")
                    f.write(f"-- {'-' * 50}\n\n")
                    
                    for cambio in self.cambios:
                        cuenta = cambio['cuenta']
                        valor_nuevo = cambio['correcto']
                        nivel = cambio['nivel']
                        
                        if valor_nuevo is None:
                            f.write(f"-- Cuenta: {cuenta} (Nivel {nivel}) | Actual: {cambio['actual_str']} -> NULL\n")
                            f.write(f"UPDATE arcgms SET cta_niv3 = NULL WHERE cuenta = {cuenta};\n\n")
                        else:
                            f.write(f"-- Cuenta: {cuenta} (Nivel {nivel}) | Actual: {cambio['actual_str']} -> {str(valor_nuevo).zfill(10)}\n")
                            f.write(f"UPDATE arcgms SET cta_niv3 = {valor_nuevo} WHERE cuenta = {cuenta};\n\n")
                    
                    f.write(f"-- TOTAL DE ACTUALIZACIONES: {len(self.cambios)}\n")
                else:
                    f.write(f"-- RESULTADO: No se requieren actualizaciones.\n")
                    f.write(f"-- Todas las cuentas tienen cta_niv3 correctamente asignado.\n")
            
            print(f"\nArchivo SQL generado: {archivo_sql}")
            return archivo_sql
            
        except Exception as e:
            print(f"Error al generar archivo SQL: {e}")
            return None
    
    def generar_reporte(self, archivo_reporte=None):
        """Genera un reporte de los cambios realizados"""
        if archivo_reporte is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            archivo_reporte = f'scripts/reporte_cta_niv3_{timestamp}.txt'
        
        try:
            with open(archivo_reporte, 'w', encoding='utf-8') as f:
                f.write(f"REPORTE DE ACTUALIZACION CTA_NIV3 - PLAN DE CUENTAS\n")
                f.write(f"{'='*60}\n")
                f.write(f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write(f"Archivo procesado: {self.archivo_excel}\n")
                f.write(f"Total de cuentas: {len(self.df)}\n")
                f.write(f"Cambios realizados: {len(self.cambios)}\n\n")
                
                if self.cambios:
                    f.write("CAMBIOS REALIZADOS:\n")
                    f.write("-" * 80 + "\n")
                    f.write(f"{'Fila':<6} {'Cuenta':<12} {'Nivel':<6} {'Anterior':<12} {'Nuevo':<12} {'Descripcion'}\n")
                    f.write("-" * 80 + "\n")
                    
                    for cambio in self.cambios:
                        f.write(f"{cambio['fila_excel']:<6} {cambio['cuenta']:<12} {cambio['nivel']:<6} "
                               f"{cambio['actual_str']:<12} {cambio['correcto_str']:<12} {cambio['descripcion']}\n")
                else:
                    f.write("No se realizaron cambios. Todas las cuentas tenian cta_niv3 correcto.\n")
            
            print(f"\nReporte guardado en: {archivo_reporte}")
            return archivo_reporte
            
        except Exception as e:
            print(f"ERROR al generar reporte: {e}")
            return None

def main():
    """Funcion principal"""
    parser = argparse.ArgumentParser(
        description='Actualizador de cta_niv3 del plan de cuentas',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ejemplos:
  python actualizar_cta_niv3.py
  python actualizar_cta_niv3.py plan_ctas/plan_cuentas.xlsx
  python actualizar_cta_niv3.py --reporte reporte_cambios.txt
        """
    )
    
    parser.add_argument('archivo', nargs='?', 
                       help='Ruta al archivo Excel del plan de cuentas')
    parser.add_argument('--reporte', '-r', 
                       help='Nombre del archivo de reporte')
    parser.add_argument('--sin-backup', action='store_true',
                       help='No crear backup del archivo original')
    parser.add_argument('--solo-analisis', action='store_true',
                       help='Solo mostrar cambios sin aplicarlos')
    
    args = parser.parse_args()
    
    print("ACTUALIZADOR DE CTA_NIV3 - PLAN DE CUENTAS SISK13")
    print("=" * 50)
    
    # Crear actualizador
    actualizador = ActualizadorCtaNiv3(args.archivo)
    
    # Cargar archivo
    if not actualizador.cargar_archivo():
        sys.exit(1)
    
    # Analizar cambios necesarios
    hay_cambios = actualizador.analizar_cambios_necesarios()
    
    # Mostrar cambios pendientes
    actualizador.mostrar_cambios_pendientes()
    
    if not hay_cambios:
        print("\nNo hay cambios que realizar. El archivo esta correcto.")
        sys.exit(0)
    
    # Si solo es analisis, generar SQL y salir
    if args.solo_analisis:
        actualizador.generar_sql_actualizacion()
        print("\nModo solo-analisis activado. No se aplicaron cambios.")
        sys.exit(0)
    
    # Confirmar cambios
    respuesta = input(f"\n¿Aplicar {len(actualizador.cambios)} cambios? (s/N): ")
    if respuesta.lower() != 's':
        print("Operacion cancelada por el usuario.")
        sys.exit(0)
    
    # Crear backup
    if not args.sin_backup:
        backup_file = actualizador.crear_backup()
        if backup_file is None:
            print("ERROR: No se pudo crear backup. Operacion cancelada.")
            sys.exit(1)
    
    # Aplicar cambios
    if not actualizador.aplicar_cambios():
        print("ERROR: No se pudieron aplicar los cambios.")
        sys.exit(1)
    
    # Guardar archivo
    if not actualizador.guardar_archivo():
        print("ERROR: No se pudo guardar el archivo.")
        sys.exit(1)
    
    # Generar archivo SQL con actualizaciones
    actualizador.generar_sql_actualizacion()
    
    # Generar reporte
    actualizador.generar_reporte(args.reporte)
    
    print(f"\nPROCESO COMPLETADO EXITOSAMENTE")
    print(f"- Cambios aplicados: {len(actualizador.cambios)}")
    print(f"- Celdas resaltadas: {len(actualizador.cambios)} (fondo amarillo)")
    print(f"- Archivo actualizado: {actualizador.archivo_excel}")
    
    sys.exit(0)

if __name__ == "__main__":
    main()